"""
API integration tests for Reports endpoints:
- GET  /api/v1/reports
- POST /api/v1/reports/preview
- POST /api/v1/reports/export
- GET  /api/v1/reports/export
- Security: RBAC 403 enforcement
- End-to-end export verification with database seeding and audit trail validation
"""

from datetime import datetime, timezone
import io
import uuid
import openpyxl
import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config.constants import JobStatus
from app.infrastructure.database.models.document_models import (
    Document,
    ProcessingJob,
)
from app.infrastructure.database.models.person_models import (
    ContactInformation,
    Person,
    Profession,
    ProfessionalCategory,
)
from app.infrastructure.database.models.resume_models import (
    Education,
    ExperienceSummary,
)
from app.infrastructure.database.models.user_models import AuditLog
from app.infrastructure.database.session import get_db_session
from app.infrastructure.ocr.mock_ocr_provider import MockOCRProvider
from app.infrastructure.storage.local_storage import LocalStorageProvider
from app.main import create_application
from app.presentation.dependencies.auth import (
    AuthenticatedUser,
    get_current_user,
)
from app.presentation.dependencies.document_dependencies import (
    get_ocr,
    get_storage,
)


@pytest_asyncio.fixture
async def reports_client(test_db: AsyncSession, tmp_path):
    app = create_application()

    app.dependency_overrides[get_db_session] = lambda: test_db

    test_storage = LocalStorageProvider(base_path=tmp_path)
    app.dependency_overrides[get_storage] = lambda: test_storage
    app.dependency_overrides[get_ocr] = lambda: MockOCRProvider()

    test_user = AuthenticatedUser(
        user_id=str(uuid.uuid4()),
        role="DIRECTOR",
        permissions=["documents:read", "documents:write"],
        jti="reports-test-jti",
    )
    app.dependency_overrides[get_current_user] = lambda: test_user

    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://test",
    ) as client:
        yield client


@pytest_asyncio.fixture
async def unauthorized_reports_client(test_db: AsyncSession):
    app = create_application()
    app.dependency_overrides[get_db_session] = lambda: test_db

    low_priv_user = AuthenticatedUser(
        user_id=str(uuid.uuid4()),
        role="CANDIDATE",
        permissions=[],
        jti="unauthorized-reports-jti",
    )
    app.dependency_overrides[get_current_user] = lambda: low_priv_user

    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://test",
    ) as client:
        yield client


@pytest.mark.asyncio
async def test_get_reports_catalog(reports_client: AsyncClient) -> None:
    resp = await reports_client.get("/api/v1/reports")
    assert resp.status_code == 200
    data = resp.json()

    assert isinstance(data, list)
    assert len(data) == 6
    report_types = [item["report_type"] for item in data]
    assert "INVENTORY" in report_types
    assert "PROFESSIONAL_CLASSIFICATION" in report_types
    assert "ACADEMIC" in report_types
    assert "EXPERIENCE" in report_types
    assert "GEOGRAPHIC" in report_types
    assert "CUSTOM_FILTERED" in report_types


@pytest.mark.asyncio
async def test_preview_report_endpoint(reports_client: AsyncClient) -> None:
    payload = {
        "report_type": "INVENTORY",
        "report_format": "EXCEL",
        "filters": {},
    }
    resp = await reports_client.post("/api/v1/reports/preview?limit=10", json=payload)
    assert resp.status_code == 200
    data = resp.json()

    assert data["report_type"] == "INVENTORY"
    assert "columns" in data
    assert len(data["columns"]) > 0
    assert "rows" in data
    assert "total_records" in data


@pytest.mark.asyncio
async def test_export_report_post_excel(reports_client: AsyncClient) -> None:
    payload = {
        "report_type": "INVENTORY",
        "report_format": "EXCEL",
        "filters": {},
    }
    resp = await reports_client.post("/api/v1/reports/export", json=payload)
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'attachment; filename="reporte_inventario_' in resp.headers["content-disposition"]
    assert len(resp.content) > 500


@pytest.mark.asyncio
async def test_export_report_post_pdf(reports_client: AsyncClient) -> None:
    payload = {
        "report_type": "INVENTORY",
        "report_format": "PDF",
        "filters": {},
    }
    resp = await reports_client.post("/api/v1/reports/export", json=payload)
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert 'attachment; filename="reporte_inventario_' in resp.headers["content-disposition"]
    assert resp.content.startswith(b"%PDF-")


@pytest.mark.asyncio
async def test_export_report_get_query_params(reports_client: AsyncClient) -> None:
    resp = await reports_client.get("/api/v1/reports/export?report_type=GEOGRAPHIC&report_format=EXCEL")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'attachment; filename="reporte_distribucion_geografica_' in resp.headers["content-disposition"]


@pytest.mark.asyncio
async def test_reports_forbidden_without_permission(
    unauthorized_reports_client: AsyncClient,
) -> None:
    resp = await unauthorized_reports_client.get("/api/v1/reports")
    assert resp.status_code == 403

    resp_exp = await unauthorized_reports_client.post(
        "/api/v1/reports/export",
        json={"report_type": "INVENTORY", "report_format": "EXCEL"},
    )
    assert resp_exp.status_code == 403


@pytest.mark.asyncio
async def test_full_report_export_with_database_records(
    reports_client: AsyncClient,
    test_db: AsyncSession,
) -> None:
    """Verify end-to-end report generation with seeded DB records and audit logging."""
    # 1. Seed database with candidate, category, profession, education, experience and document
    cat = ProfessionalCategory(
        id=uuid.uuid4(),
        name="Ciencias Sociales y Humanidades",
        description="Área social",
    )
    test_db.add(cat)
    await test_db.flush()

    prof = Profession(
        id=uuid.uuid4(),
        category_id=cat.id,
        name="Psicólogo Organizacional",
    )
    test_db.add(prof)
    await test_db.flush()

    person = Person(
        id=uuid.uuid4(),
        first_name="Beatriz",
        first_surname="Valencia",
        identification_type="CC",
        identification_number="43526178",
        primary_category_id=cat.id,
        primary_profession_id=prof.id,
    )
    test_db.add(person)
    await test_db.flush()

    contact = ContactInformation(
        id=uuid.uuid4(),
        person_id=person.id,
        department="Santander",
        municipality="Bucaramanga",
        email="beatriz.valencia@example.org",
    )
    test_db.add(contact)

    edu = Education(
        id=uuid.uuid4(),
        person_id=person.id,
        level="UNDERGRADUATE",
        degree_title="Psicología",
        institution="Universidad Industrial de Santander",
    )
    test_db.add(edu)

    exp = ExperienceSummary(
        id=uuid.uuid4(),
        person_id=person.id,
        total_years=6,
        private_years=6,
    )
    test_db.add(exp)

    doc = Document(
        id=uuid.uuid4(),
        person_id=person.id,
        uploaded_by=uuid.uuid4(),
        original_filename="cv_beatriz_valencia.pdf",
        storage_key="documents/cv_beatriz_valencia.pdf",
        checksum_sha256="abcdef1234567890abcdef1234567890abcdef1234567890abcdef1234567890",
        file_size_bytes=314572,
        mime_type="application/pdf",
        document_type="ATS",
        page_count=2,
    )
    test_db.add(doc)
    await test_db.flush()

    job = ProcessingJob(
        id=uuid.uuid4(),
        document_id=doc.id,
        status=JobStatus.COMPLETED.value,
    )
    test_db.add(job)
    await test_db.commit()

    # 2. Export Professional Classification Excel
    resp_excel = await reports_client.post(
        "/api/v1/reports/export",
        json={
            "report_type": "PROFESSIONAL_CLASSIFICATION",
            "report_format": "EXCEL",
            "filters": {"min_years_experience": 5},
        },
    )
    assert resp_excel.status_code == 200
    excel_bytes = resp_excel.content
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    # Search for candidate in worksheet
    found_candidate = False
    for row in ws.iter_rows(values_only=True):
        if row and any("Beatriz Valencia" in str(cell) for cell in row if cell is not None):
            found_candidate = True
            break
    assert found_candidate, "Candidate 'Beatriz Valencia' should be in the exported Excel spreadsheet"

    # 3. Export PDF Report
    resp_pdf = await reports_client.post(
        "/api/v1/reports/export",
        json={
            "report_type": "PROFESSIONAL_CLASSIFICATION",
            "report_format": "PDF",
            "filters": {},
        },
    )
    assert resp_pdf.status_code == 200
    assert resp_pdf.content.startswith(b"%PDF-")

    # 4. Verify that AuditLog captured the export event
    audit_stmt = select(AuditLog).where(AuditLog.action == "EXPORT_REPORT")
    audit_res = await test_db.execute(audit_stmt)
    logs = audit_res.scalars().all()
    assert len(logs) >= 2
    assert any(log.details.get("report_type") == "PROFESSIONAL_CLASSIFICATION" for log in logs)
