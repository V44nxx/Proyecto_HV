"""
Unit tests for ExcelReportGenerator and PdfReportGenerator.
"""

from datetime import datetime, timezone
import io
import openpyxl
import pytest

from app.domain.value_objects.report_types import (
    ReportColumn,
    ReportDataset,
    ReportType,
)
from app.infrastructure.reporting.excel_report_generator import (
    ExcelReportGenerator,
)
from app.infrastructure.reporting.pdf_report_generator import (
    PdfReportGenerator,
)


@pytest.fixture
def sample_dataset() -> ReportDataset:
    now = datetime(2026, 9, 20, 15, 30, 0, tzinfo=timezone.utc)
    columns = [
        ReportColumn("number", "N°", width_excel=8, width_pdf_ratio=0.5, align="center"),
        ReportColumn("candidate_name", "Candidato", width_excel=25, width_pdf_ratio=2.0),
        ReportColumn("profession", "Profesión", width_excel=25, width_pdf_ratio=2.0),
        ReportColumn("experience_years", "Años Exp.", width_excel=14, width_pdf_ratio=1.0, align="right"),
        ReportColumn("active", "Activo", width_excel=10, width_pdf_ratio=0.8, align="center"),
    ]
    rows = [
        {
            "number": 1,
            "candidate_name": "Ana María Gómez",
            "profession": "Ingeniera de Sistemas",
            "experience_years": 7,
            "active": True,
        },
        {
            "number": 2,
            "candidate_name": "Juan Pablo Rodríguez",
            "profession": "Contador Público",
            "experience_years": 12,
            "active": False,
        },
    ]
    return ReportDataset(
        report_type=ReportType.PROFESSIONAL_CLASSIFICATION,
        title="Reporte Institucional de Clasificación",
        description="Muestra representativa de profesionales y trayectoria para pruebas unitarias.",
        generated_at=now,
        generated_by="test_admin@institucion.gov.co",
        applied_filters={"min_years_experience": 5},
        columns=columns,
        rows=rows,
        total_records=2,
    )


def test_excel_report_generator(sample_dataset: ReportDataset):
    generator = ExcelReportGenerator()
    content = generator.generate(sample_dataset)

    assert isinstance(content, bytes)
    assert len(content) > 1000

    # Validate workbook integrity with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert "Reporte" in wb.sheetnames
    ws = wb["Reporte"]

    # Title and description
    assert ws["A1"].value == sample_dataset.title
    assert ws["A2"].value == sample_dataset.description

    # Header and Data check
    # Row 4 is metadata
    # Row 9 is header row
    found_header = False
    for row in ws.iter_rows(values_only=True):
        if row and "Candidato" in row:
            found_header = True
            break
    assert found_header, "Header row with 'Candidato' not found in worksheet"

    # Verify data exists in worksheet
    found_candidate = False
    for row in ws.iter_rows(values_only=True):
        if row and "Ana María Gómez" in row:
            found_candidate = True
            break
    assert found_candidate, "Candidate row not found in worksheet"


def test_pdf_report_generator(sample_dataset: ReportDataset):
    generator = PdfReportGenerator()
    content = generator.generate(sample_dataset)

    assert isinstance(content, bytes)
    assert len(content) > 1000
    # Standard PDF magic header
    assert content.startswith(b"%PDF-")


def test_pdf_report_generator_large_dataset():
    """Verify that multi-page tables render with NumberedCanvas without page overflow errors."""
    now = datetime(2026, 9, 20, 12, 0, 0, tzinfo=timezone.utc)
    columns = [
        ReportColumn("number", "N°", width_excel=8, width_pdf_ratio=0.5, align="center"),
        ReportColumn("title", "Título Largo de Cargo y Responsabilidades", width_excel=40, width_pdf_ratio=3.0),
        ReportColumn("company", "Empresa u Organización Extensa", width_excel=30, width_pdf_ratio=2.5),
        ReportColumn("years", "Años", width_excel=10, width_pdf_ratio=0.8, align="right"),
    ]
    rows = [
        {
            "number": i,
            "title": f"Cargo Especializado de Consultoría y Dirección de Proyectos Tecnológicos Número {i}",
            "company": f"Corporación Internacional de Desarrollo y Soluciones Globales {i} S.A.S.",
            "years": i % 15,
        }
        for i in range(1, 80)  # Generates multiple pages
    ]
    dataset = ReportDataset(
        report_type=ReportType.EXPERIENCE,
        title="Reporte Extenso de Trayectoria Multicriterio",
        description="Dataset con más de 70 filas para validar paginación automática y numeración.",
        generated_at=now,
        generated_by="audit@test.com",
        applied_filters={"limit": 80},
        columns=columns,
        rows=rows,
        total_records=len(rows),
    )

    generator = PdfReportGenerator()
    content = generator.generate(dataset)

    assert isinstance(content, bytes)
    assert content.startswith(b"%PDF-")
    assert len(content) > 5000  # Multi-page PDF size
