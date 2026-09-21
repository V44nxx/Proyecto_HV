"""
Excel report generator using openpyxl.
Produces professionally styled spreadsheets with institutional styling and metadata.
"""

import io
from typing import Any
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.domain.value_objects.report_types import ReportDataset


class ExcelReportGenerator:
    """Renders a ReportDataset into an Excel (.xlsx) file in bytes."""

    # Institutional color palette
    HEADER_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    TITLE_FONT = Font(name="Calibri", size=15, bold=True, color="1B365D")
    SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="475569")
    META_LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="334155")
    META_VAL_FONT = Font(name="Calibri", size=10, color="334155")
    BODY_FONT = Font(name="Calibri", size=10, color="0F172A")

    THIN_BORDER = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    def generate(self, dataset: ReportDataset) -> bytes:
        """Generates the .xlsx file content for the given dataset."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte"
        ws.views.sheetView[0].showGridLines = True

        # 1. Institutional Title and Description
        ws.cell(row=1, column=1, value=dataset.title).font = self.TITLE_FONT
        ws.cell(row=2, column=1, value=dataset.description).font = self.SUBTITLE_FONT

        # 2. Metadata Block
        date_str = dataset.generated_at.strftime("%Y-%m-%d %H:%M:%S UTC")
        filters_str = self._format_filters(dataset.applied_filters)

        meta_rows = [
            ("Fecha de Generación:", date_str),
            ("Generado por:", dataset.generated_by),
            ("Filtros aplicados:", filters_str),
            ("Total de Registros:", str(dataset.total_records)),
        ]

        curr_row = 4
        for label, val in meta_rows:
            lbl_cell = ws.cell(row=curr_row, column=1, value=label)
            lbl_cell.font = self.META_LABEL_FONT
            val_cell = ws.cell(row=curr_row, column=2, value=val)
            val_cell.font = self.META_VAL_FONT
            curr_row += 1

        curr_row += 1  # Empty spacer line

        # 3. Table Column Headers
        header_row_idx = curr_row
        ws.row_dimensions[header_row_idx].height = 24

        col_alignments = {}
        for col_idx, col in enumerate(dataset.columns, start=1):
            cell = ws.cell(row=header_row_idx, column=col_idx, value=col.label)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            align_horizontal = col.align or "left"
            cell.alignment = Alignment(horizontal=align_horizontal, vertical="center", wrap_text=True)
            cell.border = self.THIN_BORDER
            col_alignments[col_idx] = align_horizontal

        curr_row += 1

        # 4. Table Data Rows
        for row_idx, record in enumerate(dataset.rows, start=curr_row):
            is_even = (row_idx % 2 == 0)
            fill = self.ZEBRA_FILL if is_even else self.WHITE_FILL
            ws.row_dimensions[row_idx].height = 20

            for col_idx, col in enumerate(dataset.columns, start=1):
                raw_val = record.get(col.key)
                display_val = self._format_cell_value(raw_val)

                cell = ws.cell(row=row_idx, column=col_idx, value=display_val)
                cell.font = self.BODY_FONT
                cell.fill = fill
                cell.border = self.THIN_BORDER
                cell.alignment = Alignment(
                    horizontal=col_alignments[col_idx],
                    vertical="center",
                )

        # 5. Column Width Optimization
        for col_idx, col in enumerate(dataset.columns, start=1):
            col_letter = get_column_letter(col_idx)
            # Find max length of contents in this column
            max_len = len(str(col.label))
            for record in dataset.rows[:100]:  # Sample first 100 rows for speed
                val = record.get(col.key)
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            # Clamp between col.width_excel and 50
            optimal_width = max(col.width_excel, min(max_len + 4, 50))
            ws.column_dimensions[col_letter].width = optimal_width

        # Write to in-memory bytes
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def _format_cell_value(self, value: Any) -> Any:
        if value is None:
            return ""
        if isinstance(value, bool):
            return "Sí" if value else "No"
        if hasattr(value, "strftime"):
            return value.strftime("%Y-%m-%d %H:%M")
        if isinstance(value, list):
            return ", ".join(str(v) for v in value if v)
        return value

    def _format_filters(self, filters: dict[str, Any]) -> str:
        if not filters:
            return "Ninguno (Catálogo completo)"
        parts = []
        for k, v in filters.items():
            if v is not None and v != "":
                parts.append(f"{k}: {v}")
        return " | ".join(parts) if parts else "Ninguno (Catálogo completo)"
